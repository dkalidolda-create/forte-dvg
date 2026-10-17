"""Offline preparation only. No AI calls, network requests or face recognition."""
import argparse
import csv
import hashlib
import html
import json
import random
import re
from collections import Counter, defaultdict
from pathlib import Path

import cv2
from openpyxl import load_workbook


def normalized_id(value):
    return str(value).strip().upper().replace('Х', 'X').replace('–', '-').replace('—', '-')


def seconds(value):
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).lower().strip()
    minute = re.search(r'(\d+)\s*мин', text)
    second = re.search(r'(\d+(?:[.,]\d+)?)\s*сек', text)
    if minute or second:
        return (int(minute[1]) * 60 if minute else 0) + (float(second[1].replace(',', '.')) if second else 0)
    raise ValueError('Unrecognized timestamp: ' + text)


def write_csv(path, rows, columns):
    with path.open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def prepare(labels, videos, output, seed=42):
    if output.exists():
        raise ValueError('Output already exists. Choose a new directory; reviewer edits must be preserved.')
    files = {}
    for path in sorted(videos.iterdir()):
        if path.suffix.lower() not in {'.mov', '.mp4'}:
            continue
        key = normalized_id(path.stem)
        if key in files:
            raise ValueError('Duplicate normalized video ID: ' + key)
        files[key] = path.resolve()
    book = load_workbook(labels, read_only=True, data_only=True)
    rows = []
    seen = set()
    try:
        for name, timestamp, label in book['labels'].iter_rows(min_row=2, max_col=3, values_only=True):
            if name is None:
                continue
            key = normalized_id(name)
            if key in seen or key not in files:
                raise ValueError('Duplicate or unmatched label ID: ' + key)
            seen.add(key)
            label = str(label).strip().lower()
            if label not in {'нарушение', 'без нарушения'}:
                raise ValueError('Unexpected label for ' + key)
            target = int(label == 'нарушение')
            event = seconds(timestamp)
            if target and event is None:
                raise ValueError('Positive video without event time: ' + key)
            if not target and event is not None:
                raise ValueError('Negative video with event time: ' + key)
            path = files[key]
            cap = cv2.VideoCapture(str(path))
            try:
                fps = cap.get(cv2.CAP_PROP_FPS)
                count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
                ok, frame = cap.read()
                if not ok or fps <= 0 or count <= 0:
                    raise ValueError('Cannot decode video: ' + key)
                duration = count / fps
                if event is not None and not 0 <= event < duration:
                    raise ValueError('Event outside video: ' + key)
                height, width = frame.shape[:2]
            finally:
                cap.release()
            rows.append({'video_id': key, 'original_filename': path.name,
                         'video_path': str(path), 'target': target,
                         'event_sec': event, 'duration_sec': round(duration, 4),
                         'fps': round(fps, 6), 'frame_count': count,
                         'width': width, 'height': height,
                         'sha256': sha256(path)})
    finally:
        book.close()
    if seen != set(files):
        raise ValueError('Videos without labels: ' + ', '.join(sorted(set(files) - seen)))

    # Exact duplicate files must stay together. Related scenes require manual grouping later.
    groups = defaultdict(list)
    for row in rows:
        groups[row['sha256']].append(row)
    by_class = defaultdict(list)
    for digest, members in groups.items():
        targets = {row['target'] for row in members}
        if len(targets) != 1:
            raise ValueError('Exact duplicate videos have conflicting labels')
        by_class[members[0]['target']].append(digest)
    rng = random.Random(seed)
    validation = set()
    for target in sorted(by_class):
        digests = sorted(by_class[target])
        rng.shuffle(digests)
        count = max(1, round(len(digests) * .2)) if len(digests) > 1 else 0
        validation.update(digests[:count])
    for row in rows:
        row['split'] = 'validation' if row['sha256'] in validation else 'train'
        row['exact_file_group'] = row['sha256'][:16]

    output.mkdir(parents=True)
    write_csv(output / 'manifest.csv', rows, list(rows[0]))
    review = []
    pages = []
    for row in rows:
        key = row['video_id']
        folder = output / 'frames' / key
        folder.mkdir(parents=True)
        event = row['event_sec']
        fps = row['fps']
        last_time = (row['frame_count'] - 1) / fps
        if event is None:
            timestamps = [row['duration_sec'] * i / 8 for i in range(1, 8)]
        else:
            timestamps = [event + offset for offset in [-6, -4, -2, 0, 2, 4, 6]]
        indexes = sorted({max(0, min(row['frame_count'] - 1, int(max(0, min(t, last_time)) * fps))) for t in timestamps})
        images = []
        cap = cv2.VideoCapture(row['video_path'])
        try:
            for number, index in enumerate(indexes):
                cap.set(cv2.CAP_PROP_POS_FRAMES, index)
                ok, frame = cap.read()
                if not ok:
                    raise ValueError(f'Cannot decode review frame: {key}, {index}')
                actual = max(0, cap.get(cv2.CAP_PROP_POS_FRAMES) - 1) / fps
                h, w = frame.shape[:2]
                scale = min(1, 960 / max(h, w))
                if scale < 1:
                    frame = cv2.resize(frame, (round(w * scale), round(h * scale)))
                filename = f'{number:02d}_{actual:.3f}s.jpg'
                # imencode/tofile works with Cyrillic paths on Windows too.
                ok, encoded = cv2.imencode('.jpg', frame, [cv2.IMWRITE_JPEG_QUALITY, 82])
                if not ok:
                    raise ValueError('JPEG encoding failed')
                encoded.tofile(str(folder / filename))
                images.append(f'<figure><img loading="lazy" src="frames/{key}/{filename}"><figcaption>{actual:.3f} s</figcaption></figure>')
        finally:
            cap.release()
        review.append({'video_id': key, 'split': row['split'], 'provided_target': row['target'],
                       'provided_event_sec': event, 'reviewer': '', 'observed_action': '',
                       'event_type': '', 'event_start_sec': '', 'event_end_sec': '',
                       'rule': '', 'source_title': '', 'source_url': '', 'source_section': '',
                       'related_scene_group': row['exact_file_group'] if len(groups[row['sha256']]) > 1 else '',
                       'review_status': 'pending', 'notes': ''})
        title = f'{key}: ' + ('нарушение' if row['target'] else 'без нарушения')
        src = html.escape(Path(row['video_path']).as_uri(), quote=True)
        start = max(0, (event or 0) - 6)
        pages.append(f'<section id="{key}"><h2>{html.escape(title)}</h2>'
                     f'<p>{row["split"]}; длительность {row["duration_sec"]:.1f} s; '
                     f'метка времени: {event if event is not None else "нет"}</p>'
                     f'<p><a href="{src}#t={start}" target="_blank">Открыть исходное видео</a></p>'
                     '<div class="grid">' + ''.join(images) + '</div></section>')
    write_csv(output / 'review.csv', review, list(review[0]))
    summary = {'videos': len(rows), 'labels': dict(Counter(row['target'] for row in rows)),
               'splits': dict(Counter(row['split'] for row in rows)), 'seed': seed,
               'duration_min_sec': min(row['duration_sec'] for row in rows),
               'duration_max_sec': max(row['duration_sec'] for row in rows),
               'exact_duplicate_groups': sum(len(members) > 1 for members in groups.values()),
               'duplicates': [
                   [{'video_id': row['video_id'], 'event_sec': row['event_sec'],
                     'target': row['target'], 'split': row['split']} for row in members]
                   for members in groups.values() if len(members) > 1
               ],
               'split_warning': 'Provisional. Check related scenes/actors and event classes before training.',
               'timing_warning': 'FPS-based seek. Verify actual timing for variable-frame-rate videos.'}
    (output / 'audit.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    nav = ' · '.join(f'<a href="#{r["video_id"]}">{r["video_id"]}</a>' for r in rows)
    document = ('<!doctype html><html lang="ru"><meta charset="utf-8"><title>Forte — разбор данных</title>'
                '<style>body{font:16px system-ui;margin:32px;background:#111827;color:#eee}'
                'a{color:#93c5fd}section{border-top:1px solid #475569;padding:20px 0}'
                '.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:12px}'
                'figure{margin:0}img{width:100%}figcaption{padding:6px}</style>'
                '<h1>Локальный разбор обучающих видео</h1>'
                '<p>Кадры не являются предсказаниями модели. Заполняйте review.csv, '
                'проверяя полное видео. Не отправляйте видео и кадры в облачные AI-сервисы.</p>'
                '<p>Для положительных видео: контекст ±6 секунд. Для отрицательных: '
                '7 обзорных кадров; они не заменяют просмотр всего ролика.</p><nav>' + nav + '</nav>'
                + ''.join(pages) + '</html>')
    (output / 'index.html').write_text(document, encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--labels', type=Path, required=True)
    parser.add_argument('--videos', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--seed', type=int, default=42)
    args = parser.parse_args()
    prepare(args.labels, args.videos, args.output, args.seed)


if __name__ == '__main__':
    main()
